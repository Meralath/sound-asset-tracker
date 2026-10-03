"""Build a colour-coded Excel tracker from a game's sound list.

Example:
    python sound_tracker.py sounds.csv --title "Forest Cabin" -o tracker.xlsx

The sound list is either a CSV file with a header row (columns: name,
category, type, source, spatial; only name is required) or a TXT file with
one sound name per line. Run with --help to see every option.
"""

import argparse
import csv
import io
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    print("error: openpyxl is not installed. Install it with: pip install -r requirements.txt",
          file=sys.stderr)
    sys.exit(1)


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

# The columns a CSV sound list can have. Only "name" is required.
INPUT_COLUMNS = ["name", "category", "type", "source", "spatial"]

# The choices in each dropdown. Values in a sound list must be one of these.
TYPE_CHOICES = ["OneShot", "Loop", "Sting"]
STATUS_CHOICES = ["Not Started", "Recording", "Editing", "Mixing", "Done"]
SOURCE_CHOICES = ["Original", "Library", "Hybrid"]
SPATIAL_CHOICES = ["2D", "3D"]

DEFAULT_TYPE = "OneShot"
DEFAULT_STATUS = "Not Started"

# An Excel cell holds at most 32,767 characters; longer text is cut off when the
# file is opened. The tracker adds a few words to the title and the section
# headers, so values from the list must stay a little under that.
MAX_TEXT_LENGTH = 32000

# The columns of the sheet and their widths.
COLUMNS = [
    ("Asset Name", 45),
    ("Type", 12),
    ("Status", 16),
    ("Source", 14),
    ("Spatialization", 16),
    ("Notes", 50),
]
LAST_COLUMN = get_column_letter(len(COLUMNS))

# Section header colour (dark) and row tint (light) for common game-audio categories.
PALETTE = {
    "MUSIC":       ("4A1A5C", "EFE5F5"),
    "ENEMY":       ("6B0F0F", "F5E0E0"),
    "PLAYER":      ("1A4A50", "DEEAEC"),
    "ITEMS":       ("8B5A00", "F5EAD0"),
    "ENVIRONMENT": ("3D3D3D", "EAEAEA"),
    "AMBIENCE":    ("1F3D2A", "DEE8E1"),
    "VO":          ("2A1F5C", "E2DEF2"),
    "UI":          ("2D3D4A", "DEE3E7"),
}

# Other spellings that get the same colour as a palette category. Categories
# often come from name prefixes, so Item_Kettle_Whistle (category ITEM) still gets
# the ITEMS colour. The category keeps the name it was given.
COLOUR_ALIASES = {
    "ITEM": "ITEMS",
    "ENV": "ENVIRONMENT",
    "AMB": "AMBIENCE",
    "ENEMIES": "ENEMY",
    "DIALOGUE": "VO",
}

# Colours for any other category. They are handed out in this order as new
# categories appear, so the same sound list always gets the same colours.
FALLBACK_COLOURS = [
    ("7A3E1D", "F6E6DC"),  # rust
    ("1D4E7A", "DDE8F2"),  # blue
    ("5C1A45", "F2DDEA"),  # plum
    ("4A5C1A", "E9EEDB"),  # olive
    ("1A5C52", "DBEEEB"),  # teal
    ("7A1D3E", "F5DDE5"),  # raspberry
    ("5A4632", "EEE8E1"),  # taupe
    ("2F6B2F", "E0EEE0"),  # green
]

TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
SUBTITLE_FONT = Font(name="Calibri", size=11, italic=True, color="BBBBBB")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Calibri", size=10)
NOTES_FONT = Font(name="Calibri", size=10, italic=True, color="888888")

CENTRE = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", indent=1)

HEADER_BORDER = Border(bottom=Side(style="thin", color="FFFFFF"))
ROW_BORDER = Border(bottom=Side(style="thin", color="DDDDDD"))


class InputError(Exception):
    """A problem with the sound list or the output path, shown as a one-line message."""


# ---------------------------------------------------------------------------
# Reading the sound list
# ---------------------------------------------------------------------------

def read_sound_list(path):
    """Read a .csv or .txt sound list.

    Returns the sounds (one dict each) and the CSV columns that were ignored.
    """
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix == ".txt":
        return read_txt(path), []
    raise InputError(f"{path} is not a .csv or .txt file")


def read_text(path):
    """Return the whole file as text."""
    try:
        # "utf-8-sig" also reads the invisible marker Excel puts at the start of UTF-8 files.
        return path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        raise InputError(f"{path} is not UTF-8 text (in Excel, save it as 'CSV UTF-8')")
    except OSError as error:
        raise InputError(f"could not read {path}: {error.strerror}")


def read_csv(path):
    """Read a CSV sound list. Returns the sounds and the names of any columns it ignored."""
    text = read_text(path)
    if not text.strip():
        raise InputError(f"{path} is empty")
    header_line = text.split("\n")[0]
    if not header_line.strip():
        raise InputError(f"{path} line 1 is blank; the first line must be the header row "
                         f"(for example: name,category,type,source,spatial)")

    rows = split_csv_rows(path, text, guess_delimiter(header_line))

    # The first row is the header. Its cells get the same check as every other row,
    # so a quote left open in the header cannot swallow the rows below it.
    _line_number, header_fields = rows[0]
    for position, value in enumerate(header_fields, start=1):
        check_cell(value, f"the header in column {position}", f"{path} line 1")

    # Accept headers such as "Name" or " name ".
    header = [field.strip() for field in header_fields]
    columns = [field.lower() for field in header]
    if "name" not in columns:
        found = ", ".join(field for field in header if field) or "nothing"
        raise InputError(f"{path} has no 'name' column in its header row (found: {found})")
    ignored = [field for field in header if field and field.lower() not in INPUT_COLUMNS]

    sounds = []
    for line_number, fields in rows[1:]:
        where = f"{path} line {line_number}"
        check_row(fields, header, where)

        # Pair each value with its column name. A short row simply has fewer values.
        row = dict(zip(columns, fields))
        values = {}
        for column in INPUT_COLUMNS:
            values[column] = row.get(column, "").strip()

        # Skip empty rows, such as the blank lines Excel sometimes leaves at the end.
        if not any(values.values()):
            continue

        if not values["name"]:
            raise InputError(f"{where}: the name is empty")
        sounds.append(make_sound(values["name"], values["category"], values["type"],
                                 values["source"], values["spatial"], where))
    return sounds, ignored


def guess_delimiter(header_line):
    """Return ";" or ",", whichever the header line has more of outside quotes.

    Excel in many European locales saves CSV files with semicolons instead of
    commas. Characters inside quotes are skipped, so a header cell such as
    "notes, extra" does not count as a comma.
    """
    commas = 0
    semicolons = 0
    inside_quotes = False
    for character in header_line:
        if character == '"':
            inside_quotes = not inside_quotes
        elif character == "," and not inside_quotes:
            commas += 1
        elif character == ";" and not inside_quotes:
            semicolons += 1
    if semicolons > commas:
        return ";"
    return ","


def find_bad_character(text):
    """Describe the first character that a one-line cell cannot hold, or return "".

    Line breaks and tabs would break the sheet's rows. Besides the usual line
    feed (code 10) and carriage return (13), three Unicode characters also
    count as line breaks: next line (133), line separator (8232) and
    paragraph separator (8233).

    Some characters cannot be stored in an Excel file at all, and saving would
    stop with an error: the other invisible control characters (codes 0 to 31),
    the codes 65534 and 65535, and the leftovers of broken text (codes 55296
    to 57343), which a title typed on Linux can contain.
    """
    for character in text:
        code = ord(character)
        if code in (10, 13, 133, 8232, 8233):
            return "a line break"
        if code == 9:
            return "a tab"
        if code < 32 or code in (65534, 65535) or 55296 <= code <= 57343:
            return f"a character that Excel cannot store (code {code})"
    return ""


def check_length(text, what, where):
    """Stop if text is too long for an Excel cell."""
    if len(text) > MAX_TEXT_LENGTH:
        raise InputError(f"{where}: {what} is {len(text):,} characters long; the limit is "
                         f"{MAX_TEXT_LENGTH:,}, because an Excel cell holds at most 32,767")


def check_cell(value, what, where):
    """Stop if a CSV cell is too long, or holds a character a one-line spreadsheet cell cannot hold."""
    check_length(value, what, where)
    problem = find_bad_character(value)
    if not problem:
        return
    if problem == "a line break":
        # A quote that is never closed also puts line breaks in a cell: the rows
        # after it are read as part of that cell and would otherwise vanish.
        hint = 'remove it from the cell (Alt+Enter in Excel adds one), or close a quote (") left open'
    else:
        hint = "remove it from the cell"
    raise InputError(f"{where}: {what} contains {problem}; {hint}")


def check_row(fields, header, where):
    """Check every value in a row, including those in columns the tracker does not use."""
    for position, value in enumerate(fields):
        if position < len(header) and header[position]:
            what = f"the value in the {header[position]} column"
        else:
            what = f"the value in column {position + 1}"
        check_cell(value, what, where)


def split_csv_rows(path, text, delimiter):
    """Split CSV text into rows. Returns [(line number where the row starts, [values])]."""
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows = []
    last_line_read = 0
    try:
        for fields in reader:
            # A row is usually one line, but a quoted value can run over several.
            rows.append((last_line_read + 1, fields))
            last_line_read = reader.line_num
    except csv.Error as error:
        # For example a value over the csv module's size limit, which a quote
        # that is never closed can produce in a long file.
        raise InputError(f"{path} line {last_line_read + 1}: could not read this row as CSV "
                         f'({error}); check for a quote (") that is never closed')
    return rows


def read_txt(path):
    """Read a TXT sound list: one name per line. Blank lines and lines starting with # are skipped."""
    sounds = []
    # Split on line feeds only: splitlines() would also split on some invisible
    # characters and quietly turn one name into two.
    for line_number, line in enumerate(read_text(path).split("\n"), start=1):
        name = line.strip()
        if not name or name.startswith("#"):
            continue
        where = f"{path} line {line_number}"
        check_length(name, "the name", where)
        # Check the line as written, so a tab or control character at either end
        # is reported instead of quietly dropped, the same as in a CSV.
        problem = find_bad_character(line)
        if problem == "a tab":
            # A tab often means the file has several columns, which only a CSV can hold.
            raise InputError(f"{where}: the name contains a tab; remove it, or save a list "
                             f"with several columns as a .csv file")
        if problem:
            raise InputError(f"{where}: the name contains {problem}; remove it")
        sounds.append(make_sound(name, where=where))
    return sounds


def make_sound(name, category="", sound_type="", source="", spatial="", where=""):
    """Fill in the defaults and check each value against the dropdown choices."""
    if category:
        category = category.upper()
    else:
        category = category_from_name(name)

    sound_type = match_choice(sound_type, TYPE_CHOICES, "type", where)
    if not sound_type:
        sound_type = DEFAULT_TYPE

    return {
        "name": name,
        "category": category,
        "type": sound_type,
        "source": match_choice(source, SOURCE_CHOICES, "source", where),
        "spatial": match_choice(spatial, SPATIAL_CHOICES, "spatial", where),
    }


def category_from_name(name):
    """Use the part of the name before the first underscore: Player_Step_01 -> PLAYER."""
    first_part = name.split("_")[0]
    # A name that starts with an underscore has no first part to use.
    if not first_part:
        return "OTHER"
    return first_part.upper()


def match_choice(value, choices, column, where):
    """Return the choice that matches value, spelled as in the dropdown ("loop" -> "Loop").

    An empty value stays empty. A value that matches no choice stops the run:
    it is usually a typo, and the dropdown would never offer it.
    """
    if not value:
        return ""
    for choice in choices:
        if simplify(choice) == simplify(value):
            return choice
    raise InputError(f"{where}: {column} '{value}' is not one of {', '.join(choices)}")


def simplify(text):
    """Lower-case the text and drop spaces and hyphens, so "One-Shot" matches "OneShot"."""
    return text.lower().replace(" ", "").replace("-", "")


def group_by_category(sounds):
    """Return {category: [sounds]}, with categories in the order they first appear."""
    groups = {}
    for sound in sounds:
        if sound["category"] not in groups:
            groups[sound["category"]] = []
        groups[sound["category"]].append(sound)
    return groups


def pick_colours(categories):
    """Take a list of categories and return {category: (dark, light)}.

    A palette category, or one of its aliases, gets its palette colour.
    Any other category gets the next fallback colour.
    """
    colours = {}
    fallbacks_used = 0
    for category in categories:
        palette_name = COLOUR_ALIASES.get(category, category)
        if palette_name in PALETTE:
            colours[category] = PALETTE[palette_name]
        else:
            # With more new categories than fallback colours, start the list again.
            colours[category] = FALLBACK_COLOURS[fallbacks_used % len(FALLBACK_COLOURS)]
            fallbacks_used += 1
    return colours


# ---------------------------------------------------------------------------
# Building the workbook
# ---------------------------------------------------------------------------

def solid_fill(hex_colour):
    return PatternFill(fill_type="solid", fgColor=hex_colour)


def plural(count, singular, plural_form):
    """plural(1, "Asset", "Assets") -> "1 Asset"; plural(5, ...) -> "5 Assets"."""
    if count == 1:
        return f"{count} {singular}"
    return f"{count} {plural_form}"


def write_text(sheet, row, column, text):
    """Put text in a cell as plain text and return the cell.

    Marking the cell as text matters: openpyxl would save text that starts
    with "=" as a formula, and Excel would show an error instead of the name.
    """
    cell = sheet.cell(row=row, column=column)
    # Leave empty cells truly empty instead of storing an empty text.
    if text:
        cell.value = text
        cell.data_type = "s"
    return cell


def write_banner(sheet, row, text, font, fill_colour, height, alignment):
    """Write one full-width row (merged across all columns)."""
    sheet.merge_cells(f"A{row}:{LAST_COLUMN}{row}")
    cell = write_text(sheet, row, 1, text)
    cell.font = font
    cell.fill = solid_fill(fill_colour)
    cell.alignment = alignment
    sheet.row_dimensions[row].height = height


def write_column_headers(sheet, row):
    for column, (label, _width) in enumerate(COLUMNS, start=1):
        cell = write_text(sheet, row, column, label)
        cell.font = HEADER_FONT
        cell.fill = solid_fill("2D2D2D")
        cell.alignment = CENTRE
        cell.border = HEADER_BORDER
    sheet.row_dimensions[row].height = 22


def write_asset_row(sheet, row, sound, fill_colour):
    """Write one sound's row. The Notes column starts empty, ready for the user."""
    values = [sound["name"], sound["type"], DEFAULT_STATUS, sound["source"], sound["spatial"], ""]
    notes_column = len(values)
    for column, value in enumerate(values, start=1):
        cell = write_text(sheet, row, column, value)
        cell.fill = solid_fill(fill_colour)
        cell.border = ROW_BORDER
        if column == notes_column:
            cell.font = NOTES_FONT
        else:
            cell.font = ROW_FONT
        # Names and notes read better left-aligned; the short dropdown values centred.
        if column in (1, notes_column):
            cell.alignment = LEFT
        else:
            cell.alignment = CENTRE


def add_dropdowns(sheet, asset_blocks):
    """Add the four dropdowns to the asset rows only, never to headers or spacer rows.

    asset_blocks holds (first_row, last_row) for each category's asset rows.
    """
    dropdown_columns = [
        ("B", TYPE_CHOICES),
        ("C", STATUS_CHOICES),
        ("D", SOURCE_CHOICES),
        ("E", SPATIAL_CHOICES),
    ]
    for column, choices in dropdown_columns:
        dropdown = DataValidation(type="list", formula1='"' + ",".join(choices) + '"',
                                  allow_blank=True)
        for first_row, last_row in asset_blocks:
            dropdown.add(f"{column}{first_row}:{column}{last_row}")
        sheet.add_data_validation(dropdown)


def build_workbook(groups, title=""):
    """Lay out the tracker: title, subtitle, column headers, then one section per category."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Asset Tracker"

    for column, (_label, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    total = sum(len(sounds) for sounds in groups.values())
    if title:
        heading = f"{title.upper()} - SOUND ASSET TRACKER"
    else:
        heading = "SOUND ASSET TRACKER"
    subtitle = plural(total, "Asset", "Assets") + " - " + plural(len(groups), "Category", "Categories")

    write_banner(sheet, 1, heading, TITLE_FONT, "000000", 32, CENTRE)
    write_banner(sheet, 2, subtitle, SUBTITLE_FONT, "1A1A1A", 20, CENTRE)
    write_column_headers(sheet, 3)
    # Keep the title, subtitle and column headers in view while scrolling.
    sheet.freeze_panes = "A4"

    colours = pick_colours(list(groups))
    asset_blocks = []
    row = 4
    for category, sounds in groups.items():
        dark, light = colours[category]
        count_text = plural(len(sounds), "asset", "assets")
        write_banner(sheet, row, f"▌ {category} ({count_text})", SECTION_FONT, dark, 24, LEFT)
        row += 1

        first_asset_row = row
        for position, sound in enumerate(sounds):
            # Alternate the category's tint with white so long lists are easy to follow.
            if position % 2 == 0:
                fill_colour = light
            else:
                fill_colour = "FFFFFF"
            write_asset_row(sheet, row, sound, fill_colour)
            row += 1
        asset_blocks.append((first_asset_row, row - 1))

        # A thin empty row between categories.
        sheet.row_dimensions[row].height = 8
        row += 1

    add_dropdowns(sheet, asset_blocks)
    return workbook


# ---------------------------------------------------------------------------
# Command line
# ---------------------------------------------------------------------------

class ArgumentParser(argparse.ArgumentParser):
    """argparse's parser with one change, to how a bad command line is reported.

    On its own, argparse prints a usage block and exits with code 2. This tool
    reports every problem the same way, as one "error:" line and exit code 1,
    so the error method is replaced here.
    """

    def error(self, message):
        print(f"error: {message} (run with --help to see the options)", file=sys.stderr)
        sys.exit(1)


def build_parser():
    parser = ArgumentParser(
        description="Build a colour-coded Excel tracker (.xlsx) from a game's sound list,\n"
                    "with one section per category and dropdowns for Type, Status,\n"
                    "Source and Spatialization.",
        epilog='example:\n  python sound_tracker.py sounds.csv --title "Forest Cabin" -o tracker.xlsx',
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "sound_list",
        help="your sound list: a .csv file with a header row (columns name, category, type, "
             "source, spatial; only name is required) or a .txt file with one sound name per line",
    )
    parser.add_argument(
        "-t", "--title",
        default="",
        help='the game or project name shown in the title row, for example "Forest Cabin"',
    )
    parser.add_argument(
        "-o", "--output",
        help="where to save the tracker; must end in .xlsx. If left out, it is saved next to "
             "the sound list with the same name (sounds.csv gives sounds.xlsx)",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="replace the output file if it already exists. Without this, an existing file "
             "is never changed",
    )
    return parser


def print_summary(sound_list_label, output_label, groups):
    total = sum(len(sounds) for sounds in groups.values())
    print(f"Read {plural(total, 'sound', 'sounds')} in "
          f"{plural(len(groups), 'category', 'categories')} from {sound_list_label}")
    for category, sounds in groups.items():
        print(f"  {category:<14}{len(sounds):>3}")
    print(f"Saved {output_label}")


def run(args):
    sound_list = Path(args.sound_list)
    if args.output:
        output = Path(args.output)
        output_label = args.output
    else:
        output = sound_list.with_suffix(".xlsx")
        output_label = str(output)

    title = args.title.strip()
    check_length(title, "the title", "--title")
    problem = find_bad_character(title)
    if problem:
        raise InputError(f"the title contains {problem}; remove it")

    if sound_list.is_dir():
        raise InputError(f"{args.sound_list} is a folder; give the path to a .csv or .txt file")
    if not sound_list.is_file():
        raise InputError(f"{args.sound_list} not found")
    sounds, ignored_columns = read_sound_list(sound_list)
    if not sounds:
        raise InputError(f"no sound names found in {args.sound_list}")

    if output.suffix.lower() != ".xlsx":
        raise InputError(f"the output file must end in .xlsx (got {output_label})")
    if output.is_dir():
        raise InputError(f"{output_label} is a folder, not a file")
    if output.exists() and not args.force:
        raise InputError(f"{output_label} already exists; use --force to replace it")
    if not output.parent.is_dir():
        raise InputError(f"the folder for {output_label} does not exist")

    groups = group_by_category(sounds)
    workbook = build_workbook(groups, title)
    try:
        workbook.save(output)
    except PermissionError:
        raise InputError(f"could not save {output_label}: permission denied (is it open in Excel?)")
    except OSError as error:
        raise InputError(f"could not save {output_label}: {error.strerror}")

    # Warn only once the tracker is saved, so a run that fails still prints a single line.
    if ignored_columns:
        if len(ignored_columns) == 1:
            label = "column"
        else:
            label = "columns"
        print(f"warning: ignored the {label} {', '.join(ignored_columns)} in {args.sound_list}; "
              f"the columns used are {', '.join(INPUT_COLUMNS)}", file=sys.stderr)
    print_summary(args.sound_list, output_label, groups)


def main(argv=None):
    args = build_parser().parse_args(argv)
    try:
        run(args)
    except InputError as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    # Print in UTF-8 so a name with letters such as Ş or İ cannot crash the output.
    # Without this, Windows writes to a file or pipe in an older code page that lacks them.
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
