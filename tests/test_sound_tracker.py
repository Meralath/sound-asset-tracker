"""Tests for sound_tracker.py.

Most tests run the script the way a user would, then open the saved
workbook and check what is actually in it.
"""

import csv
import os
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

import sound_tracker

TOOL_FOLDER = Path(__file__).resolve().parent.parent
SCRIPT = TOOL_FOLDER / "sound_tracker.py"
EXAMPLE_CSV = TOOL_FOLDER / "examples" / "sounds_example.csv"
TXT_FIXTURE = Path(__file__).resolve().parent / "fixtures" / "three_sounds.txt"

SECTION_HEADER = re.compile(r"^▌ (.+) \((\d+) assets?\)$")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def run_tool(*args):
    """Run the script as a user would. Returns (exit code, normal output, error output)."""
    command = [sys.executable, str(SCRIPT)] + [str(arg) for arg in args]
    result = subprocess.run(command, capture_output=True, text=True,
                            encoding="utf-8", errors="replace")
    return result.returncode, result.stdout, result.stderr


def assert_one_line_error(error_output):
    lines = error_output.strip().splitlines()
    assert len(lines) == 1, error_output
    assert lines[0].startswith("error: ")
    assert "Traceback" not in error_output


def read_tracker(path):
    """Read a saved tracker back into a plain summary the tests can check."""
    sheet = openpyxl.load_workbook(path).active
    tracker = {
        "sheet": sheet,
        "title": sheet["A1"].value,
        "subtitle": sheet["A2"].value,
        "headers": [sheet.cell(row=3, column=column).value for column in range(1, 7)],
        "sections": [],    # [category, count shown in its header, [its asset rows]]
        "asset_rows": [],
        "dropdowns": {},   # column letter -> (choices, rows it covers)
        "dropdown_count": len(sheet.data_validations.dataValidation),
    }

    for row in range(4, sheet.max_row + 1):
        value = sheet.cell(row=row, column=1).value
        header = SECTION_HEADER.match(value or "")
        if header:
            tracker["sections"].append([header.group(1), int(header.group(2)), []])
        elif value:
            tracker["sections"][-1][2].append(row)
            tracker["asset_rows"].append(row)

    for dropdown in sheet.data_validations.dataValidation:
        cells = set()
        for cell_range in dropdown.sqref.ranges:
            for row in range(cell_range.min_row, cell_range.max_row + 1):
                for column in range(cell_range.min_col, cell_range.max_col + 1):
                    cells.add((get_column_letter(column), row))
        columns = {column for column, _row in cells}
        assert len(columns) == 1, "each dropdown should cover a single column"
        choices = dropdown.formula1.strip('"').split(",")
        tracker["dropdowns"][columns.pop()] = (choices, {row for _column, row in cells})

    return tracker


def section_counts(tracker):
    """{category: (count shown in its header, number of rows under it)}"""
    return {category: (shown, len(rows)) for category, shown, rows in tracker["sections"]}


def row_values(tracker, name):
    """The six cell values of the asset row with this name."""
    sheet = tracker["sheet"]
    for row in tracker["asset_rows"]:
        if sheet.cell(row=row, column=1).value == name:
            return [sheet.cell(row=row, column=column).value for column in range(1, 7)]
    raise AssertionError(f"{name} is not in the tracker")


def colour_of(cell):
    """The cell's fill colour as six hex digits (openpyxl adds an alpha prefix)."""
    return cell.fill.fgColor.rgb[-6:]


def header_colours(tracker):
    """{category: colour of its section header}"""
    sheet = tracker["sheet"]
    colours = {}
    for category, _count, rows in tracker["sections"]:
        colours[category] = colour_of(sheet.cell(row=rows[0] - 1, column=1))
    return colours


def expected_from_csv(path):
    """Count sounds per category straight from the CSV, following the README's rules."""
    counts = {}
    with open(path, newline="", encoding="utf-8") as file:
        for row in csv.DictReader(file):
            category = row["category"].strip().upper()
            if not category:
                category = row["name"].split("_")[0].upper()
            counts[category] = counts.get(category, 0) + 1
    return counts


@pytest.fixture
def example_tracker(tmp_path):
    output = tmp_path / "tracker.xlsx"
    code, _out, error = run_tool(EXAMPLE_CSV, "--title", "Forest Cabin", "-o", output)
    assert code == 0, error
    return read_tracker(output)


# ---------------------------------------------------------------------------
# The example sound list
# ---------------------------------------------------------------------------

def test_example_counts_match_the_csv(example_tracker):
    expected = expected_from_csv(EXAMPLE_CSV)
    total = sum(expected.values())
    assert total == 22  # the example itself has 22 sounds in 7 categories
    assert len(expected) == 7

    assert len(example_tracker["asset_rows"]) == total
    assert example_tracker["subtitle"] == f"{total} Assets - {len(expected)} Categories"

    found = section_counts(example_tracker)
    assert list(found) == list(expected)  # same categories, in the order they first appear
    for category, count in expected.items():
        assert found[category] == (count, count)


def test_every_dropdown_covers_exactly_the_asset_rows(example_tracker):
    asset_rows = set(example_tracker["asset_rows"])
    assert example_tracker["dropdown_count"] == 4
    assert example_tracker["dropdowns"] == {
        "B": (["OneShot", "Loop", "Sting"], asset_rows),
        "C": (["Not Started", "Recording", "Editing", "Mixing", "Done"], asset_rows),
        "D": (["Original", "Library", "Hybrid"], asset_rows),
        "E": (["2D", "3D"], asset_rows),
    }


def test_example_layout(example_tracker):
    sheet = example_tracker["sheet"]
    assert example_tracker["title"] == "FOREST CABIN - SOUND ASSET TRACKER"
    assert example_tracker["headers"] == [
        "Asset Name", "Type", "Status", "Source", "Spatialization", "Notes"]
    assert sheet.freeze_panes == "A4"

    merged = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert {"A1:F1", "A2:F2"} <= merged
    for _category, _count, rows in example_tracker["sections"]:
        header_row = rows[0] - 1
        assert f"A{header_row}:F{header_row}" in merged


def test_example_cell_values(example_tracker):
    # Blank type becomes OneShot; every sound starts as Not Started; Notes is empty.
    assert row_values(example_tracker, "Player_Step_Leaves_01") == [
        "Player_Step_Leaves_01", "OneShot", "Not Started", "Library", "2D", None]
    assert row_values(example_tracker, "Music_Nightfall_Sting") == [
        "Music_Nightfall_Sting", "Sting", "Not Started", "Original", "2D", None]

    # A category given in the CSV wins over the name's first part.
    ambience_rows = [rows for category, _count, rows in example_tracker["sections"]
                     if category == "AMBIENCE"][0]
    names = [example_tracker["sheet"].cell(row=row, column=1).value for row in ambience_rows]
    assert "Cabin_Fireplace_Crackle_Loop" in names


def test_palette_colours_and_fallback_colours(example_tracker):
    colours = header_colours(example_tracker)
    for category in ("MUSIC", "PLAYER", "AMBIENCE", "UI"):
        assert colours[category] == sound_tracker.PALETTE[category][0]
    # ITEM comes from the prefix of Item_Kettle_Whistle and gets the ITEMS colour.
    assert colours["ITEM"] == sound_tracker.PALETTE["ITEMS"][0]
    # Categories outside the palette take the fallback colours in order of appearance.
    assert colours["WILDLIFE"] == sound_tracker.FALLBACK_COLOURS[0][0]
    assert colours["WEATHER"] == sound_tracker.FALLBACK_COLOURS[1][0]


def test_rows_alternate_tint_and_white(example_tracker):
    sheet = example_tracker["sheet"]
    categories = [category for category, _count, _rows in example_tracker["sections"]]
    colours = sound_tracker.pick_colours(categories)
    for category, _count, rows in example_tracker["sections"]:
        light = colours[category][1]
        for position, row in enumerate(rows):
            expected = light if position % 2 == 0 else "FFFFFF"
            for column in range(1, 7):
                assert colour_of(sheet.cell(row=row, column=column)) == expected


def test_same_input_always_gives_the_same_colours(tmp_path):
    colour_lists = []
    for run_number in (1, 2):
        output = tmp_path / f"tracker_{run_number}.xlsx"
        code, _out, error = run_tool(EXAMPLE_CSV, "-o", output)
        assert code == 0, error
        sheet = openpyxl.load_workbook(output).active
        colour_lists.append([colour_of(sheet.cell(row=row, column=1))
                             for row in range(1, sheet.max_row + 1)])
    assert colour_lists[0] == colour_lists[1]


def test_fallback_colours_are_handed_out_in_order():
    fallback = sound_tracker.FALLBACK_COLOURS
    colours = sound_tracker.pick_colours(["ZEBRA", "MUSIC", "ALPHA"])
    assert colours == {"ZEBRA": fallback[0], "MUSIC": sound_tracker.PALETTE["MUSIC"],
                       "ALPHA": fallback[1]}

    # With more new categories than fallback colours, the list starts again.
    many = [f"CATEGORY{number}" for number in range(len(fallback) + 1)]
    colours = sound_tracker.pick_colours(many)
    assert colours[many[-1]] == fallback[0]


def test_palette_uses_generic_names_and_aliases():
    palette = sound_tracker.PALETTE
    assert list(palette) == ["MUSIC", "ENEMY", "PLAYER", "ITEMS", "ENVIRONMENT",
                             "AMBIENCE", "VO", "UI"]
    # The same eight colours the tracker has always used.
    assert sorted(palette.values()) == sorted([
        ("4A1A5C", "EFE5F5"), ("6B0F0F", "F5E0E0"), ("1A4A50", "DEEAEC"),
        ("8B5A00", "F5EAD0"), ("3D3D3D", "EAEAEA"), ("1F3D2A", "DEE8E1"),
        ("2A1F5C", "E2DEF2"), ("2D3D4A", "DEE3E7"),
    ])
    # Short prefixes such as Item_ and Env_ get the colour of the full category name.
    colours = sound_tracker.pick_colours(["ITEM", "ENV", "AMB", "ENEMIES", "DIALOGUE"])
    assert colours == {
        "ITEM": palette["ITEMS"],
        "ENV": palette["ENVIRONMENT"],
        "AMB": palette["AMBIENCE"],
        "ENEMIES": palette["ENEMY"],
        "DIALOGUE": palette["VO"],
    }


# ---------------------------------------------------------------------------
# TXT lists and the rules for filling in blanks
# ---------------------------------------------------------------------------

def test_txt_list_takes_categories_from_the_names(tmp_path):
    output = tmp_path / "tracker.xlsx"
    code, _out, error = run_tool(TXT_FIXTURE, "-o", output)
    assert code == 0, error

    tracker = read_tracker(output)
    assert section_counts(tracker) == {"PLAYER": (2, 2), "UI": (1, 1)}
    assert tracker["subtitle"] == "3 Assets - 2 Categories"
    for row in tracker["asset_rows"]:
        assert tracker["sheet"].cell(row=row, column=2).value == "OneShot"


def test_category_from_name():
    assert sound_tracker.category_from_name("Player_Footstep_01") == "PLAYER"
    assert sound_tracker.category_from_name("ui_click") == "UI"
    assert sound_tracker.category_from_name("Thunder") == "THUNDER"
    assert sound_tracker.category_from_name("_Temp_01") == "OTHER"


def test_values_are_matched_without_caring_about_case(tmp_path):
    sounds = tmp_path / "sounds.csv"
    # Written with the byte-order mark Excel adds, and with an empty row at the end.
    sounds.write_text(
        "Name,Category,Type,Source,Spatial\n"
        "Rain_Loop,ambience,loop,LIBRARY,2d\n"
        "Door_Knock,,One-Shot,hybrid,3d\n"
        ",,,,\n",
        encoding="utf-8-sig",
    )
    output = tmp_path / "tracker.xlsx"
    code, _out, error = run_tool(sounds, "-o", output)
    assert code == 0, error

    tracker = read_tracker(output)
    assert section_counts(tracker) == {"AMBIENCE": (1, 1), "DOOR": (1, 1)}
    assert row_values(tracker, "Rain_Loop") == [
        "Rain_Loop", "Loop", "Not Started", "Library", "2D", None]
    assert row_values(tracker, "Door_Knock") == [
        "Door_Knock", "OneShot", "Not Started", "Hybrid", "3D", None]


def test_semicolon_separated_csv(tmp_path):
    sounds = tmp_path / "sounds.csv"
    sounds.write_text("name;category;type\nRain_Loop;AMBIENCE;Loop\nOwl_Hoot;;\n",
                      encoding="utf-8")
    output = tmp_path / "tracker.xlsx"
    code, _out, error = run_tool(sounds, "-o", output)
    assert code == 0, error
    assert section_counts(read_tracker(output)) == {"AMBIENCE": (1, 1), "OWL": (1, 1)}


def test_one_sound_uses_singular_words(tmp_path):
    sounds = tmp_path / "sounds.txt"
    sounds.write_text("Thunder\n", encoding="utf-8")
    output = tmp_path / "tracker.xlsx"
    code, _out, error = run_tool(sounds, "-o", output)
    assert code == 0, error

    tracker = read_tracker(output)
    assert tracker["subtitle"] == "1 Asset - 1 Category"
    assert tracker["sheet"]["A4"].value == "▌ THUNDER (1 asset)"
    assert tracker["title"] == "SOUND ASSET TRACKER"  # no --title given


def test_text_that_looks_like_a_formula_stays_text(tmp_path):
    sounds = tmp_path / "sounds.txt"
    sounds.write_text('=1+1_Odd_Name\n=HYPERLINK("x")\n', encoding="utf-8")
    output = tmp_path / "tracker.xlsx"
    code, _out, error = run_tool(sounds, "--title", "=2+2", "-o", output)
    assert code == 0, error

    tracker = read_tracker(output)
    sheet = tracker["sheet"]
    first_name = sheet.cell(row=tracker["asset_rows"][0], column=1)
    assert first_name.value == "=1+1_Odd_Name"
    assert first_name.data_type == "s"  # "s" is text; a formula would be "f"
    assert sheet["A1"].value == "=2+2 - SOUND ASSET TRACKER"
    assert sheet["A1"].data_type == "s"

    # No cell anywhere in the sheet is a formula.
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.data_type != "f", cell.coordinate


def test_unknown_columns_give_one_warning_and_the_tracker_is_still_built(tmp_path):
    sounds = tmp_path / "sounds.csv"
    sounds.write_text("name,Spatialization,Status,,type\nRain_Loop,3D,Done,,Loop\n",
                      encoding="utf-8")
    output = tmp_path / "tracker.xlsx"
    code, out, error = run_tool(sounds, "-o", output)
    assert code == 0, error

    assert error.strip().splitlines() == [
        f"warning: ignored the columns Spatialization, Status in {sounds}; "
        "the columns used are name, category, type, source, spatial"]
    assert out.splitlines()[-1] == f"Saved {output}"
    assert row_values(read_tracker(output), "Rain_Loop") == [
        "Rain_Loop", "Loop", "Not Started", None, None, None]


# ---------------------------------------------------------------------------
# Command line behaviour
# ---------------------------------------------------------------------------

def test_console_summary(tmp_path):
    output = tmp_path / "tracker.xlsx"
    code, out, error = run_tool(EXAMPLE_CSV, "-o", output)
    assert code == 0, error
    assert error == ""  # no warnings for the example
    lines = out.splitlines()
    assert lines[0] == f"Read 22 sounds in 7 categories from {EXAMPLE_CSV}"
    assert lines[1].split() == ["MUSIC", "2"]
    assert lines[2].split() == ["PLAYER", "5"]
    assert lines[-1] == f"Saved {output}"


@pytest.mark.parametrize("output_encoding", [None, "cp1252"], ids=["default", "cp1252"])
def test_non_english_letters_print_without_crashing(tmp_path, output_encoding):
    # When output goes to a pipe or a file, Windows uses an older code page such
    # as cp1252, which has no Ş, İ or Ğ. The "cp1252" case forces that anywhere.
    env = dict(os.environ)
    env.pop("PYTHONIOENCODING", None)
    env.pop("PYTHONUTF8", None)
    if output_encoding:
        env["PYTHONIOENCODING"] = output_encoding

    sounds = tmp_path / "sounds.txt"
    sounds.write_text("Şehir_Rüzgar_01\nİç_Kapı_01\nĞ_Test_01\n", encoding="utf-8")
    output = tmp_path / "tracker.xlsx"
    result = subprocess.run([sys.executable, str(SCRIPT), str(sounds), "-o", str(output)],
                            capture_output=True, env=env)
    error = result.stderr.decode("utf-8", errors="replace")
    assert result.returncode == 0, error
    assert "Traceback" not in error
    out = result.stdout.decode("utf-8")
    for category in ("ŞEHIR", "İÇ", "Ğ"):
        assert category in out

    # An error message with such letters is also one clean line.
    result = subprocess.run([sys.executable, str(SCRIPT), str(tmp_path / "Şarkı.csv")],
                            capture_output=True, env=env)
    assert result.returncode == 1
    assert_one_line_error(result.stderr.decode("utf-8"))
    assert "Şarkı.csv not found" in result.stderr.decode("utf-8")


def test_default_output_is_next_to_the_sound_list(tmp_path):
    sounds = tmp_path / "sounds.txt"
    sounds.write_text(TXT_FIXTURE.read_text(encoding="utf-8"), encoding="utf-8")
    code, _out, error = run_tool(sounds)
    assert code == 0, error
    assert (tmp_path / "sounds.xlsx").is_file()


def test_existing_output_is_left_unchanged(tmp_path):
    output = tmp_path / "tracker.xlsx"
    output.write_bytes(b"my own spreadsheet")
    modified_before = output.stat().st_mtime_ns

    code, out, error = run_tool(EXAMPLE_CSV, "-o", output)
    assert code == 1
    assert out == ""
    assert_one_line_error(error)
    assert "already exists" in error
    assert output.read_bytes() == b"my own spreadsheet"
    assert output.stat().st_mtime_ns == modified_before


def test_force_replaces_an_existing_output(tmp_path):
    output = tmp_path / "tracker.xlsx"
    output.write_bytes(b"old file")
    code, _out, error = run_tool(EXAMPLE_CSV, "-o", output, "--force")
    assert code == 0, error
    assert len(read_tracker(output)["asset_rows"]) == 22


def test_csv_without_a_name_column(tmp_path):
    sounds = tmp_path / "sounds.csv"
    sounds.write_text("title,category\nPlayer_Step_01,PLAYER\n", encoding="utf-8")
    output = tmp_path / "tracker.xlsx"

    code, _out, error = run_tool(sounds, "-o", output)
    assert code == 1
    assert_one_line_error(error)
    assert "no 'name' column" in error
    assert not output.exists()


# A quote opened on line 2 and never closed, followed by enough rows to go past
# the csv module's limit of 131072 characters for one value.
LONG_FILE_WITH_AN_OPEN_QUOTE = (
    'name,category\nRain_Loop,"AMBIENCE\n'
    + "".join(f"Sound_{number:05d},WEATHER\n" for number in range(10000))
).encode("utf-8")

BAD_SOUND_LISTS = [
    # file name, file contents, part of the expected message
    pytest.param("sounds.csv", b"", "is empty", id="empty"),
    pytest.param("sounds.csv", b"name,type\nRain_Loop,Burst\n",
                 "line 2: type 'Burst' is not one of OneShot, Loop, Sting", id="bad type"),
    pytest.param("sounds.csv", b"name,source\nRain_Loop,Foley\n",
                 "source 'Foley' is not one of", id="bad source"),
    pytest.param("sounds.csv", b"name,spatial\nRain_Loop,5.1\n",
                 "spatial '5.1' is not one of", id="bad spatial"),
    pytest.param("sounds.csv", b"name,category\nRain_Loop,AMBIENCE\n,AMBIENCE\n",
                 "line 3: the name is empty", id="empty name"),
    pytest.param("sounds.csv", "name\nCafé_Ambience\n".encode("latin-1"),
                 "is not UTF-8", id="not utf-8"),
    pytest.param("sounds.txt", b"\n# only a comment\n\n",
                 "no sound names found", id="no names"),
    pytest.param("sounds.xlsx", b"not a sound list",
                 "is not a .csv or .txt file", id="wrong extension"),
    pytest.param("sounds.csv", b"name\n" + b"A" * 140000 + b"\n",
                 "line 2: could not read this row as CSV", id="value too long"),
    pytest.param("sounds.csv", LONG_FILE_WITH_AN_OPEN_QUOTE,
                 "line 2: could not read this row as CSV", id="open quote, long file"),
    pytest.param("sounds.csv", b"\nname\nRain_Loop\n",
                 "line 1 is blank; the first line must be the header row", id="blank first line"),
    pytest.param("sounds.csv", b"   \nname\nRain_Loop\n",
                 "line 1 is blank; the first line must be the header row",
                 id="first line only spaces"),

    # Line breaks and tabs: from a cell (Alt+Enter in Excel) or from a quote left open.
    pytest.param("sounds.csv",
                 b'name,category\nRain_Loop,AMBIENCE\n"Owl_Hoot,AMBIENCE\nWind_Loop,AMBIENCE\n',
                 "line 3: the value in the name column contains a line break; remove it from "
                 'the cell (Alt+Enter in Excel adds one), or close a quote (") left open',
                 id="open quote, short file"),
    pytest.param("sounds.csv", b'name\n"Two\nLines"\n',
                 "line 2: the value in the name column contains a line break; remove it",
                 id="line break in name"),
    pytest.param("sounds.csv", b'name,category\nRain_Loop,"AMB\nIENCE"\n',
                 "line 2: the value in the category column contains a line break",
                 id="line break in category"),
    pytest.param("sounds.csv", b'name,notes\nRain_Loop,"left open\nOwl_Hoot,quiet\n',
                 "line 2: the value in the notes column contains a line break",
                 id="open quote in an unused column"),
    pytest.param("sounds.csv", b'name\nRain_Loop,"extra\nvalue"\n',
                 "line 2: the value in column 2 contains a line break",
                 id="open quote past the header"),
    pytest.param("sounds.csv", b"name,category\nRain\tLoop,AMBIENCE\n",
                 "line 2: the value in the name column contains a tab; remove it from the cell",
                 id="tab in csv name"),
    # The Unicode line separator (UTF-8 bytes E2 80 A8) and next line (C2 85).
    pytest.param("sounds.csv", b"name,type\nRain_Loop,Loop\xe2\x80\xa8Owl\n",
                 "line 2: the value in the type column contains a line break",
                 id="unicode line separator in csv"),
    pytest.param("sounds.txt", b"Rain\xc2\x85Loop\n",
                 "line 1: the name contains a line break; remove it", id="next line in txt"),
    pytest.param("sounds.txt", b"Rain_Loop\nPlayer_Step_01\tPLAYER\n",
                 "line 2: the name contains a tab; remove it, or save a list with several "
                 "columns as a .csv file", id="tab in txt"),

    # The header row gets the same checks as every other row.
    pytest.param("sounds.csv", b'name,"notes\nRain_Loop,quiet\nOwl_Hoot,far\n',
                 "line 1: the header in column 2 contains a line break", id="open quote in header"),
    pytest.param("sounds.csv", b'name,"notes\nRain_Loop,quiet\nOwl_Hoot,far"\nWind_Loop,soft\n',
                 "line 1: the header in column 2 contains a line break",
                 id="open quote in header, closed later"),
    pytest.param("sounds.csv", b'name,"Spatial\nization"\nRain_Loop,3D\n',
                 "line 1: the header in column 2 contains a line break",
                 id="line break in a header cell"),
    pytest.param("sounds.csv", b"na\x01me\nRain_Loop\n",
                 "line 1: the header in column 1 contains a character that Excel cannot store "
                 "(code 1)", id="control character in header"),

    # Invisible control characters, which an Excel file cannot store.
    pytest.param("sounds.csv", b"name\nRain\x01Loop\n",
                 "line 2: the value in the name column contains a character that Excel cannot "
                 "store (code 1); remove it from the cell", id="control character in csv name"),
    pytest.param("sounds.csv", b"name,category\nRain_Loop,AMB\x0cIENCE\n",
                 "line 2: the value in the category column contains a character that Excel "
                 "cannot store (code 12)", id="control character in csv category"),
    # Python 3.10 reports a NUL itself ("line contains NUL"); later versions pass it on
    # to the tool's own check. Either way it is one line that names line 2.
    pytest.param("sounds.csv", b"name\nRain\x00Loop\n", "line 2: ", id="nul in csv name"),
    pytest.param("sounds.txt", b"Rain_Loop\nOwl\x1bHoot\n",
                 "line 2: the name contains a character that Excel cannot store (code 27); "
                 "remove it", id="control character in txt"),
    pytest.param("sounds.txt", b"Rain_Loop\t\n",
                 "line 1: the name contains a tab", id="tab at the end of a txt line"),
    pytest.param("sounds.txt", b"\x1fRain_Loop\n",
                 "line 1: the name contains a character that Excel cannot store (code 31)",
                 id="control character at the start of a txt line"),
    pytest.param("sounds.txt", b"Rain\x1cLoop\n",
                 "line 1: the name contains a character that Excel cannot store (code 28)",
                 id="separator character in txt is not split into two names"),

    # Text longer than an Excel cell can hold would be cut off when the file is opened.
    pytest.param("sounds.csv", b"name\n" + b"A" * 32001 + b"\n",
                 "line 2: the value in the name column is 32,001 characters long; "
                 "the limit is 32,000", id="csv value too long for a cell"),
    pytest.param("sounds.txt", b"Rain_Loop\n" + b"B" * 40000 + b"\n",
                 "line 2: the name is 40,000 characters long; the limit is 32,000",
                 id="txt name too long for a cell"),
]


@pytest.mark.parametrize("file_name, contents, message", BAD_SOUND_LISTS)
def test_bad_sound_lists_give_a_one_line_error(tmp_path, file_name, contents, message):
    sounds = tmp_path / file_name
    sounds.write_bytes(contents)
    output = tmp_path / "tracker.xlsx"

    code, _out, error = run_tool(sounds, "-o", output)
    assert code == 1
    assert_one_line_error(error)
    assert message in error
    assert not output.exists()


def test_bad_paths_give_a_one_line_error(tmp_path):
    folder_named_like_a_file = tmp_path / "folder.xlsx"
    folder_named_like_a_file.mkdir()
    cases = [
        ([tmp_path / "missing.csv"], "not found"),
        ([tmp_path], "is a folder; give the path to a .csv or .txt file"),
        ([EXAMPLE_CSV, "-o", tmp_path / "tracker.csv"], "must end in .xlsx"),
        ([EXAMPLE_CSV, "-o", tmp_path / "no_such_folder" / "tracker.xlsx"], "does not exist"),
        ([EXAMPLE_CSV, "-o", folder_named_like_a_file], "is a folder, not a file"),
        ([], "the following arguments are required"),
    ]
    for args, message in cases:
        code, _out, error = run_tool(*args)
        assert code == 1, args
        assert_one_line_error(error)
        assert message in error


def test_a_title_with_a_control_character_gives_a_one_line_error(tmp_path):
    output = tmp_path / "tracker.xlsx"
    code, _out, error = run_tool(EXAMPLE_CSV, "--title", "Forest\x07Cabin", "-o", output)
    assert code == 1
    assert_one_line_error(error)
    assert "the title contains a character that Excel cannot store (code 7); remove it" in error
    assert not output.exists()


@pytest.mark.skipif(sys.platform == "win32",
                    reason="only Linux and macOS pass undecodable bytes to a program")
def test_a_title_with_broken_text_gives_a_one_line_error(tmp_path):
    output = tmp_path / "tracker.xlsx"
    # The byte 0xFF is not valid UTF-8; Python hands it to the script as code 56575.
    command = [os.fsencode(sys.executable), os.fsencode(SCRIPT), os.fsencode(EXAMPLE_CSV),
               b"--title", b"Forest \xff", b"-o", os.fsencode(output)]
    result = subprocess.run(command, capture_output=True)
    error = result.stderr.decode("utf-8")
    assert result.returncode == 1
    assert_one_line_error(error)
    assert "the title contains a character that Excel cannot store (code 56575)" in error
    assert not output.exists()


def test_text_length_limit(tmp_path, capsys):
    # Run in the same process: a 32,001-character title is too long for a
    # Windows command line.
    output = tmp_path / "tracker.xlsx"
    code = sound_tracker.main([str(EXAMPLE_CSV), "--title", "T" * 32001, "-o", str(output)])
    error = capsys.readouterr().err
    assert code == 1
    assert_one_line_error(error)
    assert "the title is 32,001 characters long; the limit is 32,000" in error
    assert not output.exists()

    # A name of exactly 32,000 characters is still accepted.
    sounds = tmp_path / "sounds.txt"
    sounds.write_text("A" * 32000 + "\n", encoding="utf-8")
    assert sound_tracker.main([str(sounds), "-o", str(output)]) == 0
    assert read_tracker(output)["sheet"]["A5"].value == "A" * 32000


def test_find_bad_character():
    find = sound_tracker.find_bad_character
    for fine in ("Rain_Loop", "Şehir_Rüzgar_01", "Owl Hoot", "", "=1+1"):
        assert find(fine) == ""
    # Line feed, carriage return, and the Unicode next line, line separator
    # and paragraph separator.
    for code in (10, 13, 133, 8232, 8233):
        assert find(f"Two{chr(code)}Lines") == "a line break"
    assert find("Rain\tLoop") == "a tab"
    for code in (0, 1, 7, 11, 27, 31, 65534, 65535, 55296, 57343):
        assert find(f"Rain{chr(code)}Loop") == f"a character that Excel cannot store (code {code})"


def test_guess_delimiter():
    guess = sound_tracker.guess_delimiter
    assert guess("name,category,type") == ","
    assert guess("name;category;type") == ";"
    assert guess("name") == ","
    # Commas or semicolons inside quotes do not count.
    assert guess('name;"notes, extra";category') == ";"
    assert guess('name,"notes; extra",category') == ","


def test_semicolon_csv_with_a_quoted_comma_in_the_header(tmp_path):
    sounds = tmp_path / "sounds.csv"
    sounds.write_text('name;"notes, extra";category\nRain_Loop;"soft, far";AMBIENCE\n',
                      encoding="utf-8")
    output = tmp_path / "tracker.xlsx"
    code, _out, error = run_tool(sounds, "-o", output)
    assert code == 0, error
    assert section_counts(read_tracker(output)) == {"AMBIENCE": (1, 1)}
    assert error.strip().startswith("warning: ignored the column notes, extra in ")


def test_no_long_dashes_in_the_sheet(example_tracker):
    em_dash = chr(8212)
    en_dash = chr(8211)
    sheet = example_tracker["sheet"]
    for row in sheet.iter_rows():
        for cell in row:
            text = str(cell.value or "")
            assert em_dash not in text and en_dash not in text, cell.coordinate


def test_help_explains_every_option():
    code, out, _error = run_tool("--help")
    assert code == 0
    for option in ("sound_list", "--title", "--output", "--force"):
        assert option in out
